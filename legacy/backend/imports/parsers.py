"""Reading a POS item-wise sales report.

Petpooja's export is the first target but nothing here is Petpooja-specific: the parser finds the date,
item, quantity and amount columns by header name, so Posist/Torqus/generic CSV all work. Direct API sync
is explicitly out of scope — availability and terms are unverified (§5c).
"""

import csv
import io
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation

from django.utils.dateparse import parse_date

from core.exceptions import BusinessRuleError

#: Header aliases, lowercased. Order within a list is preference order.
COLUMNS = {
    "date": ["date", "order date", "bill date", "invoice date", "txn date"],
    "item": ["item", "item name", "product", "product name", "dish", "name", "description"],
    "quantity": ["quantity", "qty", "qty sold", "plates", "count", "units"],
    "amount": ["amount", "total", "net amount", "net sales", "revenue", "value", "gross amount"],
}


@dataclass
class ParsedRow:
    row_number: int
    date: date
    item: str
    quantity: int
    amount: Decimal


def _normalise(header: str) -> str:
    return (header or "").strip().lower().replace("_", " ").replace(".", "")


def _map_columns(headers: list[str]) -> dict[str, int]:
    normalised = [_normalise(h) for h in headers]
    mapping: dict[str, int] = {}
    for key, aliases in COLUMNS.items():
        for alias in aliases:
            if alias in normalised:
                mapping[key] = normalised.index(alias)
                break
    missing = [key for key in ("date", "item", "quantity") if key not in mapping]
    if missing:
        raise BusinessRuleError(
            "Could not find these columns in the file: "
            + ", ".join(missing)
            + f". Found: {', '.join(h for h in headers if h)}. "
            "The report needs a date, an item name and a quantity.",
            code="missing_columns",
            detail={"missing": missing, "headers": headers},
        )
    return mapping


def _decimal(value) -> Decimal:
    if value is None or str(value).strip() == "":
        return Decimal("0")
    text = str(value).replace(",", "").replace("₹", "").strip()
    try:
        return Decimal(text)
    except InvalidOperation:
        return Decimal("0")


def _int(value) -> int:
    quantity = _decimal(value)
    return int(quantity)


def _rows_from_csv(data: bytes) -> list[list]:
    text = data.decode("utf-8-sig", errors="replace")
    # Petpooja's exports vary between comma and semicolon separators depending on locale.
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def _rows_from_xlsx(data: bytes) -> list[list]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def parse_sales_file(filename: str, data: bytes) -> tuple[list[ParsedRow], list[dict]]:
    """Returns (rows, problems). A problem is a row we could not read, reported rather than dropped."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        raw = _rows_from_xlsx(data)
    elif filename.lower().endswith((".csv", ".txt")):
        raw = _rows_from_csv(data)
    else:
        raise BusinessRuleError(
            "Upload a .csv or .xlsx export from your POS.", code="unsupported_file"
        )

    raw = [row for row in raw if any(cell not in (None, "") for cell in row)]
    if not raw:
        raise BusinessRuleError("That file has no rows.", code="empty_file")

    # The header is not always the first line — Petpooja prefixes a title and a date range.
    header_index, mapping = None, None
    for index, row in enumerate(raw[:20]):
        try:
            mapping = _map_columns([str(cell) if cell is not None else "" for cell in row])
            header_index = index
            break
        except BusinessRuleError:
            continue
    if mapping is None:
        _map_columns(
            [str(cell) if cell is not None else "" for cell in raw[0]]
        )  # re-raise with detail

    rows, problems = [], []
    for offset, row in enumerate(raw[header_index + 1 :], start=header_index + 2):
        item = str(_cell(row, mapping, "item") or "").strip()
        if not item:
            continue

        raw_date = _cell(row, mapping, "date")
        day = _coerce_date(raw_date)
        if day is None:
            problems.append(
                {"row": offset, "item": item, "reason": f"unreadable date: {raw_date!r}"}
            )
            continue

        quantity = _int(_cell(row, mapping, "quantity"))
        if quantity <= 0:
            problems.append({"row": offset, "item": item, "reason": "quantity is zero or negative"})
            continue

        rows.append(
            ParsedRow(
                row_number=offset,
                date=day,
                item=item,
                quantity=quantity,
                amount=_decimal(_cell(row, mapping, "amount")),
            )
        )
    return rows, problems


def _cell(row: list, mapping: dict[str, int], key: str):
    """One cell by logical column name. Short rows are common in POS exports, so a missing trailing
    column reads as empty rather than raising."""
    index = mapping.get(key)
    if index is None or index >= len(row):
        return None
    return row[index]


def _coerce_date(value) -> date | None:
    import datetime

    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    parsed = parse_date(text[:10])
    if parsed:
        return parsed
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d %b %Y", "%d-%b-%Y", "%d/%m/%y"):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None
